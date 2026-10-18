"""Generate Architecture & Process Flow Diagrams and Populate the Prescribed Pitch Deck Template.

Uses the exact template from:
C:\\Users\\Bhuvanani\\Downloads\\Prototype Submission Deck _ Gen AI Academy APAC Edition.pptx
and saves to:
- submission/SomnoCare_Pitch_Deck.pptx
- SomnoCare_Pitch_Deck.pptx
"""

from __future__ import annotations

import os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SUBMISSION_DIR = ROOT / "submission"
SHOTS_DIR = SUBMISSION_DIR / "screenshots"
DIAGRAMS_DIR = SUBMISSION_DIR / "diagrams"
TEMPLATE_PATH = Path(r"C:\Users\Bhuvanani\Downloads\Prototype Submission Deck _ Gen AI Academy APAC Edition.pptx")

SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
DIAGRAMS_DIR.mkdir(parents=True, exist_ok=True)


def _get_font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for c in candidates:
        if os.path.exists(c):
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


def draw_process_flow_diagram(out_path: Path) -> None:
    """Create a crisp 1600x780 process flow diagram for Slide 6."""
    img = Image.new("RGB", (1600, 780), (248, 250, 252))
    draw = ImageDraw.Draw(img)

    title_f = _get_font(30, True)
    box_title_f = _get_font(21, True)
    box_body_f = _get_font(16, False)
    badge_f = _get_font(15, True)

    draw.text((48, 26), "SomnoCare End-to-End Patient & Clinical Workflow (5 Simple Steps)", fill=(15, 23, 42), font=title_f)

    steps = [
        (
            "1. Secure Sign-In &\nPersonal Profile",
            "• PBKDF2-SHA256 Auth\n• 4-Digit Privacy PIN\n• Age, Height, Weight\n• Automatic BMI & Neck\n• Smoking Quit Date",
            (2, 132, 199),
            (224, 242, 254),
        ),
        (
            "2. Clinical OSA Risk\nScreening",
            "• 8-Item STOP-BANG\n  (Low / Mod / High)\n• 8-Item Epworth (ESS)\n• Sleep Study PDF/OCR\n• Red-Flag Safety Guard",
            (79, 70, 229),
            (224, 231, 255),
        ),
        (
            "3. Nightly Sleep,\nVoice & Snore Log",
            "• WebAudio FFT Snore\n  (80–350 Hz On-Device)\n• Multilingual Voice\n  (EN / HI / KN / TE)\n• Position & Alcohol Log",
            (13, 148, 136),
            (204, 251, 241),
        ),
        (
            "4. OpenCV Pose, Yoga\n& Throat Therapy",
            "• MediaPipe 2D Skeleton\n• Joint-Angle Rep Count\n• Squats, Bridge, Cat-Cow\n• Oropharyngeal Tongue\n• 4-7-8 Pranayama",
            (5, 150, 105),
            (209, 250, 229),
        ),
        (
            "5. Gemini ADK Coach\n& Physician Report",
            "• 12 ADK Clinical Tools\n• Weight↔AHI & Snore r\n• Weather Humidifier Tune\n• Rich A2UI v0.8 Cards\n• 1-Click Doctor PDF",
            (217, 119, 6),
            (254, 243, 199),
        ),
    ]

    x = 45
    y = 105
    w = 278
    h = 370
    gap = 30

    for idx, (stitle, sbody, border_col, bg_col) in enumerate(steps):
        bx = x + idx * (w + gap)
        draw.rounded_rectangle([bx, y, bx + w, y + h], radius=22, fill=bg_col, outline=border_col, width=4)
        draw.rounded_rectangle([bx + 16, y + 16, bx + w - 16, y + 86], radius=14, fill=border_col)
        draw.text((bx + 28, y + 25), stitle, fill=(255, 255, 255), font=box_title_f)
        draw.text((bx + 24, y + 106), sbody, fill=(15, 23, 42), font=box_body_f, spacing=9)

        if idx < len(steps) - 1:
            ax1 = bx + w + 4
            ax2 = bx + w + gap - 4
            ay = y + h // 2
            draw.line([ax1, ay, ax2, ay], fill=(71, 85, 105), width=5)
            draw.polygon([(ax2, ay), (ax2 - 12, ay - 9), (ax2 - 12, ay + 9)], fill=(71, 85, 105))

    # Bottom Continuous Feedback & Persistence Loop Bar
    by1, by2 = 525, 735
    draw.rounded_rectangle([45, by1, 1555, by2], radius=22, fill=(15, 23, 42), outline=(2, 132, 199), width=3)
    draw.text(
        (75, by1 + 22),
        "Continuous Personalization, Safety & Cloud Persistence Loop (Firebase + Cloud Firestore + Cloud Run)",
        fill=(56, 189, 248),
        font=box_title_f,
    )
    bottom_items = [
        "🔒 Privacy & Consent First: On-device WebAudio FFT & MediaPipe Pose ensure raw bedroom audio/video never leave the browser.",
        "🧠 Gemini 2.5 Flash + ADK Memory Bank: Correlates BMI/Weight trends (r=+0.94) and Snoring (r=+0.97) with AHI across sessions.",
        "🚨 Clinical Red-Flag Safety Guard: Automatically halts coaching and escalates to emergency/physician care on chest pain or AHI >= 30.",
        "📄 Insurance & Doctor Ready: Real-time CMS 4h/70% compliance check + downloadable ReportLab PDF for sleep specialists.",
    ]
    for i, line in enumerate(bottom_items):
        draw.text((75, by1 + 66 + i * 32), line, fill=(241, 245, 249), font=badge_f)

    img.save(out_path)


def draw_architecture_diagram(out_path: Path) -> None:
    """Create a crisp 1600x780 system architecture diagram for Slide 8."""
    img = Image.new("RGB", (1600, 780), (248, 250, 252))
    draw = ImageDraw.Draw(img)

    title_f = _get_font(28, True)
    col_f = _get_font(20, True)
    item_t = _get_font(17, True)
    item_b = _get_font(14, False)

    draw.text((45, 22), "SomnoCare Cloud-Native Architecture (Google ADK + Gemini 2.5 Flash + Cloud Run + Firestore)", fill=(15, 23, 42), font=title_f)

    columns = [
        (
            "1. Client Edge Layer\n(Light/Dark Simplified PWA)",
            (2, 132, 199),
            (240, 249, 255),
            [
                ("5-Step Layperson UI", "Light & Dark mode, 1-click Demo Login, 4-digit PIN Lock Screen"),
                ("On-Device MediaPipe / OpenCV", "2D Pose skeleton & joint-angle rep counter (Squats, Yoga, Throat)"),
                ("On-Device WebAudio FFT", "80–350 Hz acoustic snore detector & dB meter (zero raw audio upload)"),
                ("Web Speech Voice Logger", "Hands-free sleep logging in English, Hindi, Kannada & Telugu"),
            ],
        ),
        (
            "2. Google Cloud Run\n(FastAPI + Security Gateway)",
            (79, 70, 229),
            (238, 242, 255),
            [
                ("PBKDF2 & Session Auth", "120k-iteration SHA256 hashing, HMAC tokens, brute-force lockout"),
                ("REST & Multimodal APIs", "/api/dashboard, /api/screening, /api/voice-parse, /api/weather"),
                ("Red-Flag Clinical Guard", "Deterministic safety rules for chest pain, syncope, drowsy driving"),
                ("ReportLab PDF Engine", "Generates clinical Physician PDF report (/api/doctor-report.pdf)"),
            ],
        ),
        (
            "3. Google ADK & Gemini\n(Agentic Reasoning Core)",
            (5, 150, 105),
            (236, 253, 245),
            [
                ("Gemini 2.5 Flash (Vertex/API)", "Multimodal reasoning for voice parsing, PSG PDF extraction & chat"),
                ("12 ADK Clinical Tools", "STOP-BANG, ESS, CMS 4h/70% check, BMI & Pearson r correlations"),
                ("A2UI v0.8 UI Generator", "Renders structured interactive summary cards inside agent chat"),
                ("Session Memory Bank", "Maintains longitudinal patient context across coaching turns"),
            ],
        ),
        (
            "4. Data & External APIs\n(Firestore + Live Telemetry)",
            (217, 119, 6),
            (255, 251, 235),
            [
                ("Google Cloud Firestore", "Collections: patient_profile, sleep_logs, exercise_logs, audit_logs"),
                ("Hybrid Local JSON Sync", "Zero-downtime persistence fallback for offline/edge resilience"),
                ("Open-Meteo Weather & AQI", "Live temp, humidity & PM2.5 to tune CPAP humidifier & tube temp"),
                ("Sleep Clinic Directory", "Geolocated sleep centers with 1-click Google Maps navigation"),
            ],
        ),
    ]

    cx = 40
    cy = 85
    cw = 360
    ch = 655
    cgap = 26

    for idx, (ctitle, border_c, fill_c, cards) in enumerate(columns):
        x0 = cx + idx * (cw + cgap)
        draw.rounded_rectangle([x0, cy, x0 + cw, cy + ch], radius=20, fill=fill_c, outline=border_c, width=3)
        draw.rounded_rectangle([x0 + 14, cy + 14, x0 + cw - 14, cy + 82], radius=14, fill=border_c)
        draw.text((x0 + 24, cy + 20), ctitle, fill=(255, 255, 255), font=col_f)

        ky = cy + 102
        for ktitle, kdesc in cards:
            draw.rounded_rectangle([x0 + 16, ky, x0 + cw - 16, ky + 124], radius=14, fill=(255, 255, 255), outline=border_c, width=2)
            draw.text((x0 + 28, ky + 14), ktitle, fill=border_c, font=item_t)
            # Wrap description
            words = kdesc.split()
            lines, cur = [], ""
            for w in words:
                if len(cur + " " + w) <= 42:
                    cur = (cur + " " + w).strip()
                else:
                    lines.append(cur)
                    cur = w
            if cur:
                lines.append(cur)
            draw.text((x0 + 28, ky + 44), "\n".join(lines[:3]), fill=(51, 65, 85), font=item_b, spacing=5)
            ky += 136

        if idx < len(columns) - 1:
            ax1 = x0 + cw + 3
            ax2 = x0 + cw + cgap - 3
            ay = cy + ch // 2
            draw.line([ax1, ay, ax2, ay], fill=(30, 41, 59), width=4)
            draw.polygon([(ax2, ay), (ax2 - 10, ay - 8), (ax2 - 10, ay + 8)], fill=(30, 41, 59))

    img.save(out_path)


def set_textbox_lines(shape, lines, font_size=11, color=(15, 23, 42), bold_prefix=True):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        if bold_prefix and ":" in line and not line.startswith("http"):
            head, tail = line.split(":", 1)
            r1 = p.add_run()
            r1.text = head + ":"
            r1.font.bold = True
            r1.font.size = Pt(font_size)
            r1.font.color.rgb = RGBColor(*color)
            r1.font.name = "Arial"
            r2 = p.add_run()
            r2.text = tail
            r2.font.bold = False
            r2.font.size = Pt(font_size)
            r2.font.color.rgb = RGBColor(30, 41, 59)
            r2.font.name = "Arial"
        else:
            r = p.add_run()
            r.text = line
            r.font.bold = idx == 0 and not bold_prefix
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(*color)
            r.font.name = "Arial"


def build_pitch_deck() -> None:
    proc_img = DIAGRAMS_DIR / "process_flow_diagram.png"
    arch_img = DIAGRAMS_DIR / "architecture_diagram.png"
    draw_process_flow_diagram(proc_img)
    draw_architecture_diagram(arch_img)

    prs = Presentation(str(TEMPLATE_PATH))

    # SLIDE 1: Participant Details
    s1 = prs.slides[0]
    set_textbox_lines(
        s1.shapes[1],
        [
            "Participant Details",
            "Participant Name: Bhuvanani Vattikonda",
            "Project Title: SomnoCare — Simplified AI CPAP Adherence, Sleep Apnea Screening & Lifestyle Coach",
            "Problem Statement: Build an intelligent, accessible healthcare & wellness agent using Google ADK, Gemini 2.5 Flash, Cloud Run & Firestore that empowers everyday patients to screen for Obstructive Sleep Apnea (OSA), maintain CPAP insurance compliance, and improve airway health.",
            "Live Prototype: https://somnocare-ai.netlify.app  |  GitHub: https://github.com/Bhuvanani14/somnocare",
        ],
        font_size=12,
        color=(15, 23, 42),
    )

    # SLIDE 2: Brief about the idea
    s2 = prs.slides[1]
    set_textbox_lines(
        s2.shapes[2],
        [
            "Project Vision: Make clinical-grade Sleep Apnea (OSA) screening, CPAP therapy adherence, and daily airway exercise effortless for any layperson.",
            "The Core Problem: Nearly 1 billion adults suffer from Obstructive Sleep Apnea, and up to 50% of prescribed CPAP users abandon therapy within Year 1 due to mask leaks, nasal dryness, complex medical jargon, and disconnected lifestyle habits (weight, smoking, posture).",
            "Our Solution — SomnoCare v2.0: A simplified, 5-step Light & Dark Mode companion built on Google ADK, Gemini 2.5 Flash, Cloud Firestore, and Cloud Run.",
            "1. Everyday Simplicity & Security: Jargon-free Daily Summary cards (Light Mode default + Dark Mode toggle), PBKDF2-SHA256 login, personalized onboarding, and a 4-digit Privacy PIN lock screen.",
            "2. Clinical Risk & Lifestyle Layer: Tracks Age, BMI trend, Smoking Quit Streak, official 8-item STOP-BANG & Epworth Sleepiness Scale (ESS), and correlates Weight/BMI & Snoring directly with AHI reduction.",
            "3. Edge AI Multimodal Coaching: On-device WebAudio FFT Snore Detector (80–350 Hz), MediaPipe/OpenCV 2D Pose Skeleton Coach for Squats, Yoga & Oropharyngeal Throat Exercises, Multilingual Voice Logging (EN/HI/KN/TE), Open-Meteo CPAP Humidifier tuning, and 1-Click Physician PDF Report.",
        ],
        font_size=11,
    )

    # SLIDE 3: Approach to Common Problem Statement & ADK Workflow
    s3 = prs.slides[2]
    s3.shapes[1].top = Inches(2.85)
    s3.shapes[1].height = Inches(2.55)
    set_textbox_lines(
        s3.shapes[1],
        [
            "Real-World Clinical Impact: Prevents CPAP therapy dropout (saving CMS 4h/70% insurance coverage) while addressing root-cause OSA risk factors (BMI, smoking inflammation, supine sleeping, and upper-airway muscle tone).",
            "Translating into a Working Solution with Google ADK & Gemini 2.5 Flash:",
            "• Agentic Tool Orchestration: Built `somnocare_agent` in Google ADK (`app/agent.py`) equipped with 12 deterministic clinical tools (`get_profile`, `update_profile`, `evaluate_osa_risk_and_screening`, `calculate_cms_adherence_compliance`, `log_sleep_therapy_entry`, `log_exercise_or_yoga_session`, `get_local_weather_for_humidifier`, etc.).",
            "• Statistical Correlation Engine: Computes live Pearson correlation (r) linking Weight change -> AHI (r = +0.94) and Snoring -> AHI (r = +0.97).",
            "• Deterministic Red-Flag Safety Guard: Intercepts chest pain, syncope, drowsy driving, or AHI >= 30 before coaching and escalates to emergency/physician care.",
        ],
        font_size=10,
    )

    # SLIDE 4: Opportunities & USP
    s4 = prs.slides[3]
    s4.shapes[2].top = Inches(2.25)
    s4.shapes[2].height = Inches(3.15)
    set_textbox_lines(
        s4.shapes[2],
        [
            "How SomnoCare Differs from Existing CPAP Manufacturer Apps (myAir / DreamMapper):",
            "• Holistic OSA Risk + Therapy (Not Just Machine Data): Manufacturer apps only show basic machine hours. SomnoCare unifies STOP-BANG & ESS clinical screening, BMI/Weight trends, Smoke-Free airway recovery, and positional therapy in one place.",
            "• Built for a Basic Layperson (Light & Dark Mode): Replaces confusing clinical tables with 5 plain-English steps, traffic-light status badges, and hands-free voice logging in English, Hindi, Kannada, and Telugu.",
            "• On-Device Computer Vision & Acoustic Privacy: Runs MediaPipe/OpenCV Pose rep-counting (Squats, Yoga Bridge, Cat-Cow, Oropharyngeal Throat exercises) and WebAudio FFT Snoring detection 100% locally in the browser—zero raw video or bedroom audio is ever uploaded.",
            "• Environmental CPAP Auto-Tuning: Uses live Open-Meteo temperature, humidity & PM2.5 AQI to recommend exact CPAP humidifier levels (1–8) and heated-tube temperatures to stop rainout and dry mouth.",
            "• Enterprise Security & Doctor Handoff: PBKDF2-SHA256 authentication, 4-digit PIN privacy lock, granular camera/mic consent toggles, Firestore audit logs, and 1-click ReportLab Physician PDF export.",
        ],
        font_size=10,
    )

    # SLIDE 5: List of Features Offered
    s5 = prs.slides[4]
    set_textbox_lines(
        s5.shapes[2],
        [
            "1. Simplified 5-Step UI (Light & Dark Mode): Clean layperson cards for Insurance Status, AHI, BMI, STOP-BANG & Smoke-Free Streak.",
            "2. Login, Personalization & Security Center: PBKDF2-HMAC-SHA256 login, account registration, 4-digit PIN lock screen, and Security Audit Log.",
            "3. Profile, BMI & Smoke-Free Streak: Tracks Age, Height, Weight, BMI category, Neck circumference, and days smoke-free.",
            "4. Clinical OSA Screening (STOP-BANG + ESS): Official Chung et al. 8-item STOP-BANG & Johns 8-item Epworth Sleepiness Scale with drowsy-driving alert.",
            "5. On-Device Acoustic Snore Detector & Voice Logging: 80–350 Hz WebAudio FFT snore classifier + Multilingual voice parser (EN/HI/KN/TE).",
            "6. OpenCV / MediaPipe Pose & Throat Exercise Coach: Live 2D skeleton & joint-angle rep counter for Squats, Push-ups, Yoga Bridge, Cat-Cow & Oropharyngeal tongue slides.",
            "7. Guided 4-7-8 Pranayama & Late-Meal Logger: Animated nasal breathing pacer and dinner-to-bedtime reflux risk tracker.",
            "8. Live Weather/AQI CPAP Humidifier Tuner: Open-Meteo telemetry prescribes CPAP humidifier level & heated tube temp.",
            "9. Sleep Study PDF Extractor & 1-Click Doctor PDF Report: Extracts Diagnostic AHI/SpO2 from PSG reports and exports a 30-day Physician PDF.",
            "10. Conversational Gemini 2.5 Flash Coach + A2UI v0.8 Cards: Interactive clinical chat with Red-Flag emergency safety escalation.",
        ],
        font_size=10,
    )

    # SLIDE 6: Process Flow Diagram
    s6 = prs.slides[5]
    set_textbox_lines(
        s6.shapes[2],
        ["End-to-End Patient Journey — From Secure Onboarding & Screening to Daily Coaching & Physician PDF Handoff:"],
        font_size=10,
    )
    s6.shapes.add_picture(str(proc_img), Inches(0.51), Inches(1.55), width=Inches(8.98))

    # SLIDE 7: Wireframes / Mock Diagrams of the Proposed Solution
    s7 = prs.slides[6]
    set_textbox_lines(
        s7.shapes[2],
        [
            "Simplified 5-Step Interface Wireframe Layout (Designed for Zero Cognitive Overload):",
            "• Top Bar: Brand Status Pill (Insurance Compliant) + Patient Summary + [Light/Dark Mode] + [Doctor PDF] + [PIN Lock] + [Account]",
            "• Step 1 (My Daily Summary): 5 Plain-Language KPI Cards + 14-Day CPAP/AHI/Snoring Chart + Weight/BMI vs. AHI Chart + Badges",
            "• Step 2 (Log Sleep & Snoring): [Left] WebAudio FFT Snore Microphone + Multilingual Voice Logger | [Right] Simple Nightly Sleep Form",
            "• Step 3 (My Body & Risk Check): [Left] Age, Height, Weight, Neck Size & Smoking Tracker | [Right] 8-Item STOP-BANG & Epworth Check",
            "• Step 4 (Easy Exercise & Breathing): OpenCV/MediaPipe Pose Canvas (Joint Angle + Rep Counter) + 4-7-8 Breathing Pacer + Weather Tuner",
            "• Step 5 (Ask AI Coach & Security): Gemini 2.5 Flash Chat with A2UI Cards + Privacy Consent Toggles + Security Audit Log + PSG Upload",
        ],
        font_size=10,
    )
    shot_login = SHOTS_DIR / "01_login_gateway_light.png"
    shot_pose = SHOTS_DIR / "06_opencv_pose_yoga_coach.png"
    if shot_login.exists() and shot_pose.exists():
        s7.shapes.add_picture(str(shot_login), Inches(0.65), Inches(3.38), width=Inches(4.2))
        s7.shapes.add_picture(str(shot_pose), Inches(5.10), Inches(3.38), width=Inches(4.2))

    # SLIDE 8: Architecture Diagram
    s8 = prs.slides[7]
    set_textbox_lines(
        s8.shapes[2],
        ["Layered Cloud-Native Architecture combining Edge Privacy (WebAudio/MediaPipe) with Google Cloud Run, ADK, Gemini 2.5 Flash & Firestore:"],
        font_size=10,
    )
    s8.shapes.add_picture(str(arch_img), Inches(0.51), Inches(1.55), width=Inches(8.98))

    # SLIDE 9: Technologies / Google Services Used
    s9 = prs.slides[8]
    s9.shapes[2].top = Inches(1.55)
    s9.shapes[2].height = Inches(3.85)
    set_textbox_lines(
        s9.shapes[2],
        [
            "1. Google Agent Development Kit (ADK) & Gemini 2.5 Flash (`google-adk`, `google-genai`):",
            "   • Powers `somnocare_agent` with 12 strongly-typed clinical tools, Session Memory Bank, multimodal sleep-study PDF parsing, multilingual voice extraction, and A2UI v0.8 structured card generation.",
            "2. Google Cloud Firestore & Firebase Ecosystem (`google-cloud-firestore`):",
            "   • Stores real-time document collections (`patient_profile`, `sleep_logs`, `exercise_logs`, `meal_logs`, `partner_observations`, `users_auth`, `security_audit_logs`) with hybrid local JSON fallback for 100% uptime.",
            "3. Google Cloud Run & Cloud Build (`Dockerfile`, `deploy_cloud_run.sh`, FastAPI/Uvicorn):",
            "   • Stateless, autoscaling container serving both the REST/Agent API and static SPA with strict HTTP security headers (`CSP`, `X-Frame-Options`, `X-Content-Type-Options`).",
            "4. Edge Computer Vision & Signal Processing (MediaPipe Pose / OpenCV Geometry + WebAudio API):",
            "   • Executes 30 FPS skeletal joint-angle rep counting and 80–350 Hz acoustic snore detection on-device for zero-latency feedback and total patient privacy.",
            "5. ReportLab Clinical PDF Engine & Open-Meteo Telemetry API:",
            "   • Generates physician-ready 30-day compliance reports and live weather/AQI CPAP humidifier prescriptions.",
        ],
        font_size=10,
    )

    # SLIDE 10: Snapshots of the Prototype
    s10 = prs.slides[9]
    set_textbox_lines(
        s10.shapes[3],
        ["Live Working Prototype Snapshots — Light Mode, Dark Mode, Snoring & Voice Log, and STOP-BANG/BMI Screening:"],
        font_size=10,
    )
    grid_shots = [
        (SHOTS_DIR / "02_daily_summary_light.png", Inches(0.55), Inches(1.58), Inches(4.3)),
        (SHOTS_DIR / "03_daily_summary_dark.png", Inches(5.10), Inches(1.58), Inches(4.3)),
        (SHOTS_DIR / "04_snoring_and_voice_log.png", Inches(0.55), Inches(3.48), Inches(4.3)),
        (SHOTS_DIR / "05_bmi_stopbang_screening.png", Inches(5.10), Inches(3.48), Inches(4.3)),
    ]
    for spath, lx, ty, wd in grid_shots:
        if spath.exists():
            s10.shapes.add_picture(str(spath), lx, ty, width=wd)

    out_sub = SUBMISSION_DIR / "SomnoCare_Pitch_Deck.pptx"
    out_root = ROOT / "SomnoCare_Pitch_Deck.pptx"
    prs.save(str(out_sub))
    prs.save(str(out_root))
    print(f"Saved Pitch Deck to {out_sub} and {out_root}")


if __name__ == "__main__":
    build_pitch_deck()
