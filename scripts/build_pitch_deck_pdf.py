"""Render the 11-Slide Prescribed Gen AI Academy Pitch Deck into a Compressed PDF (< 5 MB).

Reads the populated `SomnoCare_Pitch_Deck.pptx` (including the official template banners,
all 11 slide texts, process flow & architecture diagrams, and prototype screenshots) and
renders an exact 16:9 widescreen (720pt x 405pt = 10in x 5.625in) PDF:
- SomnoCare_Pitch_Deck.pdf
- submission/SomnoCare_Pitch_Deck.pdf
"""

from __future__ import annotations

import io
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

ROOT = Path(__file__).resolve().parent.parent
PPTX_PATH = ROOT / "SomnoCare_Pitch_Deck.pptx"
OUT_ROOT_PDF = ROOT / "SomnoCare_Pitch_Deck.pdf"
OUT_SUB_PDF = ROOT / "submission" / "SomnoCare_Pitch_Deck.pdf"

PAGE_W = 720.0  # 10.0 inches * 72 pt/in
PAGE_H = 405.0  # 5.625 inches * 72 pt/in


def _compress_image_for_pdf(blob: bytes, max_width: int = 1280, quality: int = 80) -> ImageReader:
    """Compress embedded image so total 11-slide PDF stays well below 2 MB (< 5 MB limit)."""
    img = Image.open(io.BytesIO(blob))
    if img.width > max_width:
        ratio = max_width / float(img.width)
        img = img.resize((max_width, int(img.height * ratio)), Image.Resampling.LANCZOS)

    out = io.BytesIO()
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        bg = Image.new("RGB", img.size, (255, 255, 255))
        rgba = img.convert("RGBA")
        bg.paste(rgba, mask=rgba.split()[3])
        bg.save(out, format="JPEG", quality=quality, optimize=True)
    else:
        img.convert("RGB").save(out, format="JPEG", quality=quality, optimize=True)
    out.seek(0)
    return ImageReader(out)


def _xml_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_pdf_from_pptx() -> None:
    prs = Presentation(str(PPTX_PATH))
    c = canvas.Canvas(str(OUT_ROOT_PDF), pagesize=(PAGE_W, PAGE_H), pageCompression=1)
    c.setTitle("SomnoCare — Gen AI Academy APAC Prototype Submission Deck")
    c.setAuthor("Bhuvanani Vattikonda & Darshan Gowda")

    for slide_idx, slide in enumerate(prs.slides, 1):
        # White slide background
        c.setFillColor(colors.HexColor("#FFFFFF"))
        c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)

        for shape in slide.shapes:
            x = shape.left.inches * 72.0
            w = shape.width.inches * 72.0
            h = shape.height.inches * 72.0
            # Convert top-down PPTX coordinates to bottom-up ReportLab coordinates
            y_top = PAGE_H - (shape.top.inches * 72.0)
            y_bottom = y_top - h

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_reader = _compress_image_for_pdf(shape.image.blob, max_width=1350, quality=82)
                c.drawImage(img_reader, x, y_bottom, width=w, height=h, preserveAspectRatio=False, mask="auto")

            elif shape.has_text_frame:
                raw_text = shape.text_frame.text.strip()
                if not raw_text:
                    continue

                paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
                if not paragraphs:
                    continue

                # Determine if this is a slide heading box or a body box
                is_heading_box = (shape.top.inches < 1.05 and len(paragraphs) <= 4 and slide_idx > 1)

                cur_y = y_top - 4.0
                for p_idx, p in enumerate(paragraphs):
                    txt = _xml_escape(p.text.strip())
                    # Format bold prefix before colon for readability
                    if ":" in txt and not txt.startswith("http") and not is_heading_box:
                        head, tail = txt.split(":", 1)
                        html_txt = f"<b>{head}:</b>{tail}"
                    elif is_heading_box or (slide_idx == 1 and p_idx == 0):
                        html_txt = f"<b>{txt}</b>"
                    else:
                        html_txt = txt

                    if is_heading_box:
                        fsize = 13.0 if p_idx == 0 else 10.5
                        leading = fsize + 3.0
                        tcolor = colors.HexColor("#0F172A")
                    elif slide_idx == 1:
                        fsize = 14.0 if p_idx == 0 else 10.5
                        leading = fsize + 4.5
                        tcolor = colors.HexColor("#0F172A")
                    else:
                        fsize = 8.8 if len(paragraphs) >= 6 else 9.5
                        leading = fsize + 3.0
                        tcolor = colors.HexColor("#1E293B")

                    style = ParagraphStyle(
                        name=f"s_{slide_idx}_{p_idx}",
                        fontName="Helvetica",
                        fontSize=fsize,
                        leading=leading,
                        textColor=tcolor,
                    )
                    para = Paragraph(html_txt, style)
                    _, used_h = para.wrap(w - 8.0, h)
                    if cur_y - used_h < 8.0:
                        break
                    para.drawOn(c, x + 4.0, cur_y - used_h)
                    cur_y -= (used_h + 3.5)

        c.showPage()

    c.save()

    import shutil
    shutil.copyfile(OUT_ROOT_PDF, OUT_SUB_PDF)
    size_kb = OUT_ROOT_PDF.stat().st_size / 1024.0
    size_mb = size_kb / 1024.0
    print(f"Saved PDF Pitch Deck: {OUT_ROOT_PDF} ({size_kb:.1f} KB / {size_mb:.2f} MB)")


if __name__ == "__main__":
    build_pdf_from_pptx()
